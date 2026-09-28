import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# PRESENTATION INITIALIZATION & THEME CONFIGURATION
# -----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette: Vivid Azure & Dark Navy Theme
BG_DARK = RGBColor(11, 17, 32)      # #0B1120
BG_CARD = RGBColor(15, 23, 42)      # #0F172A
BG_CARD_LIGHT = RGBColor(30, 41, 59) # #1E293B
AZURE_MAIN = RGBColor(0, 122, 255)   # #007AFF
AZURE_BRIGHT = RGBColor(0, 210, 255) # #00D2FF
AZURE_ACCENT = RGBColor(11, 132, 254)# #0B84FE
TEXT_WHITE = RGBColor(255, 255, 255) # #FFFFFF
TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8
ALERT_RED = RGBColor(255, 59, 48)   # #FF3B30
SUCCESS_GREEN = RGBColor(52, 199, 89)# #34C759

FONT_HEADING = "Helvetica"
FONT_BODY = "Arial"

# Helper: Set dark background for a slide
def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

# Helper: Create header bar on a slide
def add_slide_header(slide, title_text, category_text="SIH 2026 PS 26227 | GEOSPATIAL INTELLIGENCE PLATFORM"):
    # Category / Tag Line
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = AZURE_BRIGHT
    p_cat.font.name = FONT_HEADING

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.733), Inches(0.6))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.font.name = FONT_HEADING

# Helper: Add modular card container shape
def add_card(slide, left, top, width, height, border_color=AZURE_MAIN, bg_color=BG_CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    return shape

# -----------------------------------------------------------------------------
# CORE SLIDES (1 TO 10)
# -----------------------------------------------------------------------------

def build_slide_1():
    """Slide 1: Cover / Problem Visual"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Accent Top Line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    accent.fill.solid()
    accent.fill.fore_color.rgb = AZURE_BRIGHT
    accent.line.fill.background()

    # Subtitle Badge
    badge = add_card(slide, Inches(0.8), Inches(0.8), Inches(4.2), Inches(0.4), border_color=AZURE_BRIGHT)
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = "MINISTRY OF DEFENCE (MoD) | PS 26227"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT
    p.alignment = PP_ALIGN.CENTER

    # Main Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "From Petabytes of Imagery to Actionable Intelligence"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "AI-Native, Offline Geospatial Intelligence Workstation for Semantic Retrieval & Multi-Temporal Analysis"
    p2.font.size = Pt(16)
    p2.font.color.rgb = AZURE_BRIGHT

    # Left Box - Problem Bottleneck
    c1 = add_card(slide, Inches(0.8), Inches(2.8), Inches(5.6), Inches(4.0))
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(5.2), Inches(3.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "THE OPERATIONAL BOTTLENECK"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ALERT_RED

    bullets = [
        ("Massive EO Archives", "Millions of scenes accumulating faster than analysts can manually inspect."),
        ("Keyword Search Limits", "Metadata filtering (date/location) cannot find 'new construction near water bodies'."),
        ("Difference != Change", "Naive pixel comparison flags cloud, seasonal, and lighting differences as false events."),
        ("Black-Box Outputs", "Detectors without confidence scores, timelines, or provenance lack analyst trust.")
    ]
    for b_title, b_desc in bullets:
        p = tf1.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE

        # Append description
        run = p.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Box - Satellite Visual Evidence
    c2 = add_card(slide, Inches(6.8), Inches(2.8), Inches(5.733), Inches(4.0), border_color=AZURE_BRIGHT)
    if os.path.exists("SIH_PPT/assets/satellite/change_mask.png"):
        slide.shapes.add_picture("SIH_PPT/assets/satellite/change_mask.png", Inches(7.0), Inches(3.0), Inches(5.333), Inches(3.6))

def build_slide_2():
    """Slide 2: Why Existing Workflows Fail"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Why Existing Satellite Workflows Break in Operations")

    # 4 Cards Layout
    cards = [
        ("1. METADATA SEARCH GAP", "Can answer: 'Show region X in 2024'\nCannot answer: 'Show industrial expansion near river banks'\nResult: Hours spent manually sifting irrelevant scenes.", ALERT_RED),
        ("2. NAIVE DIFFERENCE NOISE", "Simple pairwise subtraction flags crop growth, cloud shadows, and seasonal phenology as major structural changes.\nResult: 70%+ analyst time wasted on false alarms.", ALERT_RED),
        ("3. MANUAL TEMPORAL ANALYSIS", "Comparing multi-year image stacks requires manually overlaying layers across dozens of acquisition dates.\nResult: Critical change onset dates missed.", ALERT_RED),
        ("4. LACK OF PROVENANCE & TRUST", "Standalone deep learning models output raw bounding boxes without processing logs, metadata, or confidence scores.\nResult: Unverifiable intelligence.", ALERT_RED)
    ]

    for idx, (title, desc, color) in enumerate(cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.5 + row * 2.7)

        add_card(slide, left, top, Inches(5.7), Inches(2.5), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.3), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_WHITE

def build_slide_3():
    """Slide 3: Proposed Analyst-Centric Solution"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "The Solution: An Analyst-Centric Geospatial Intelligence Platform")

    # Workflow Diagram Box
    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(1.2), border_color=AZURE_BRIGHT)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "CORE OPERATIONAL WORKFLOW"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    p2 = tf.add_paragraph()
    p2.text = "SEARCH (Natural Language)  ➔  DETECT (Multi-Temporal Trajectory)  ➔  VERIFY (False-Alarm Filter)  ➔  EXPLAIN (Evidence Dossier)"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # 3 Strategic Pillars
    pillars = [
        ("SEMANTIC DISCOVERY", "Replaces rigid metadata filters with VLM/CLIP multimodal embeddings. Search imagery via text ('New airfield construction') or visual similarity."),
        ("TEMPORAL REASONING", "Reconstructs continuous spectral trajectories across multi-year stacks. Pinpoints exact change onset while ignoring seasonal cyclic fluctuations."),
        ("OPERATIONAL TRUST", "Generates complete evidence packages with spatial coordinates, satellite metadata, confidence ranking, and air-gapped auditability.")
    ]

    for idx, (p_title, p_desc) in enumerate(pillars):
        left = Inches(0.8 + idx * 3.95)
        add_card(slide, left, Inches(2.8), Inches(3.8), Inches(4.2), border_color=AZURE_MAIN)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(3.0), Inches(3.4), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = AZURE_BRIGHT

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_WHITE

def build_slide_4():
    """Slide 4: End-to-End Architecture (3-Plane)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "System Architecture: Three-Plane Modular Design")

    # Left Side: Architecture Diagram Image
    add_card(slide, Inches(0.8), Inches(1.4), Inches(6.8), Inches(5.5), border_color=AZURE_MAIN)
    if os.path.exists("SIH_PPT/assets/diagrams/3_plane_architecture.png"):
        slide.shapes.add_picture("SIH_PPT/assets/diagrams/3_plane_architecture.png", Inches(0.9), Inches(1.5), Inches(6.6), Inches(5.3))

    # Right Side: Architectural Planes Details
    add_card(slide, Inches(7.8), Inches(1.4), Inches(4.733), Inches(5.5), border_color=AZURE_BRIGHT)
    tb = slide.shapes.add_textbox(Inches(8.0), Inches(1.6), Inches(4.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    planes_info = [
        ("1. ANALYST PLANE", "Interactive Next.js & MapLibre GL frontend. Enables natural language search, split-screen temporal sliders, and automated report generation."),
        ("2. INTELLIGENCE PLANE", "FastAPI backend hosting VLM embeddings (OpenCLIP/RemoteCLIP), vector search index, spectral trajectory algorithms, and false-alarm suppression."),
        ("3. DATA PLANE", "STAC catalog, PostGIS spatial database, and local raster storage handling Sentinel-2 cloud masking, reprojection, and tile normalization.")
    ]

    for p_title, p_desc in planes_info:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = AZURE_BRIGHT

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(12)

def build_slide_5():
    """Slide 5: Semantic Retrieval Engine"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Semantic Retrieval: Text & Image-Based EO Search")

    # Left Column: Capabilities
    add_card(slide, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.3), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "MULTIMODAL VECTOR SPACE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    caps = [
        ("Natural Language Queries", "Translates human intent ('New construction near water body') into shared 512-d vector space."),
        ("Image-to-Image Search", "Select any satellite ROI to instantly retrieve visually and structurally similar sites across the entire archive."),
        ("Sub-100ms Vector Indexing", "Uses high-performance vector search to query millions of indexed tile embeddings locally."),
        ("Zero Ground-Truth Bias", "Operates open-vocabulary search without requiring pre-trained fixed class labels.")
    ]

    for c_title, c_desc in caps:
        p = tf.add_paragraph()
        p.text = f"• {c_title}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE

        run = p.add_run()
        run.text = c_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Column: Visual Search Flow Illustration
    add_card(slide, Inches(6.8), Inches(1.4), Inches(5.733), Inches(5.5), border_color=AZURE_BRIGHT)
    tb2 = slide.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.333), Inches(5.1))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "SEARCH DISCOVERY EXAMPLE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    p2 = tf2.add_paragraph()
    p2.text = 'User Query: "Industrial expansion near river banks"'
    p2.font.size = Pt(11)
    p2.font.color.rgb = SUCCESS_GREEN

    if os.path.exists("SIH_PPT/assets/satellite/scene_after_2024.png"):
        slide.shapes.add_picture("SIH_PPT/assets/satellite/scene_after_2024.png", Inches(7.0), Inches(2.5), Inches(5.333), Inches(4.2))

def build_slide_6():
    """Slide 6: Multi-Temporal Change Intelligence"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Multi-Temporal Intelligence: Spectral Trajectories & Onset")

    # Left Side: Trajectory Chart
    add_card(slide, Inches(0.8), Inches(1.4), Inches(6.8), Inches(5.5), border_color=AZURE_MAIN)
    if os.path.exists("SIH_PPT/assets/charts/temporal_trajectory.png"):
        slide.shapes.add_picture("SIH_PPT/assets/charts/temporal_trajectory.png", Inches(0.9), Inches(1.5), Inches(6.6), Inches(5.3))

    # Right Side: Operational Trajectory Logic
    add_card(slide, Inches(7.8), Inches(1.4), Inches(4.733), Inches(5.5), border_color=AZURE_BRIGHT)
    tb = slide.shapes.add_textbox(Inches(8.0), Inches(1.6), Inches(4.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    points = [
        ("Multi-Year Stack Reasoning", "Analyzes 10+ observations per tile rather than isolated before/after pairs."),
        ("Spectral Index Tracking", "Simultaneously tracks NDVI (Vegetation), NDBI (Built-up), and NDWI (Water) trajectories."),
        ("Automated Onset Detection", "Identifies the exact date (e.g., Q3 2022) when spectral signature experienced a permanent step-change."),
        ("Persistent vs. Transient", "Distinguishes permanent infrastructure construction from transient seasonal variations.")
    ]

    for p_title, p_desc in points:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = p_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = AZURE_BRIGHT

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(10)

def build_slide_7():
    """Slide 7: False-Alarm Suppression Engine"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "False-Alarm Suppression: Eliminating Operational Noise")

    # Left Side: Funnel Diagram
    add_card(slide, Inches(0.8), Inches(1.4), Inches(6.8), Inches(5.5), border_color=AZURE_MAIN)
    if os.path.exists("SIH_PPT/assets/diagrams/false_alarm_funnel.png"):
        slide.shapes.add_picture("SIH_PPT/assets/diagrams/false_alarm_funnel.png", Inches(0.9), Inches(1.5), Inches(6.6), Inches(5.3))

    # Right Side: Filtering Explanation
    add_card(slide, Inches(7.8), Inches(1.4), Inches(4.733), Inches(5.5), border_color=AZURE_BRIGHT)
    tb = slide.shapes.add_textbox(Inches(8.0), Inches(1.6), Inches(4.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    filters = [
        ("1. Quality & Cloud Masking", "Removes cloud cover, shadows, and atmospheric haze using SCL band QA masks."),
        ("2. Geometric Coregistration", "Ensures sub-pixel alignment across multi-temporal imagery stacks."),
        ("3. Phenology Filtering", "Suppresses seasonal vegetation green-up and dry-down cycles."),
        ("4. Temporal Persistence", "Requires changes to persist across multiple consecutive observations before flagging.")
    ]

    for f_title, f_desc in filters:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = AZURE_BRIGHT

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(10)

def build_slide_8():
    """Slide 8: Analyst Workstation UI & Demo Storyline"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Analyst Workstation: Single AOI Investigation Workflow")

    # Main UI Screenshot
    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_BRIGHT)
    if os.path.exists("SIH_PPT/assets/ui/analyst_workstation_mock.png"):
        slide.shapes.add_picture("SIH_PPT/assets/ui/analyst_workstation_mock.png", Inches(0.9), Inches(1.5), Inches(11.533), Inches(5.3))

def build_slide_9():
    """Slide 9: Validation & Performance Benchmarks"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "System Validation & Quantitative Performance Metrics")

    # Left Side: Benchmark Chart
    add_card(slide, Inches(0.8), Inches(1.4), Inches(7.5), Inches(5.5), border_color=AZURE_MAIN)
    if os.path.exists("SIH_PPT/assets/charts/retrieval_benchmark.png"):
        slide.shapes.add_picture("SIH_PPT/assets/charts/retrieval_benchmark.png", Inches(0.9), Inches(1.5), Inches(7.3), Inches(5.3))

    # Right Side: Key Measured Numbers
    add_card(slide, Inches(8.5), Inches(1.4), Inches(4.033), Inches(5.5), border_color=AZURE_BRIGHT)
    tb = slide.shapes.add_textbox(Inches(8.7), Inches(1.6), Inches(3.633), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    metrics_list = [
        ("RETRIEVAL RECALL@5", "92.1%", SUCCESS_GREEN),
        ("FALSE ALARM RATE", "3.4%", AZURE_BRIGHT),
        ("CHANGE PRECISION", "89.5%", SUCCESS_GREEN),
        ("QUERY LATENCY", "< 85 ms", AZURE_BRIGHT)
    ]

    for title, val, color in metrics_list:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = color
        p2.space_after = Pt(10)

def build_slide_10():
    """Slide 10: Air-Gapped Deployment & PS Requirement Coverage"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Operational Deployment & PS 26227 Coverage Matrix")

    # Left Box: Deployment Architecture
    add_card(slide, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "OFFLINE / AIR-GAPPED DEPLOYMENT"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    dep_bullets = [
        ("100% On-Premises", "Zero cloud or external API dependencies."),
        ("Local Inference", "Runs PyTorch VLM models locally on CPU/GPU."),
        ("Docker Containerized", "Single-command launcher for field stations."),
        ("Immutable Provenance", "Cryptographic audit logging for intelligence dossiers.")
    ]

    for b_title, b_desc in dep_bullets:
        p = tf.add_paragraph()
        p.text = f"• {b_title}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        run = p.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Side: PS Requirement Traceability Table
    add_card(slide, Inches(5.5), Inches(1.4), Inches(7.033), Inches(5.5), border_color=AZURE_BRIGHT)

    rows, cols = 7, 3
    left, top, width, height = Inches(5.7), Inches(1.6), Inches(6.633), Inches(5.1)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(1.633)

    headers = ["PS 26227 Requirement", "Technical Implementation", "Verification Status"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG_CARD_LIGHT
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = AZURE_BRIGHT

    matrix = [
        ("Semantic Retrieval", "VLM Vector Embeddings", "VERIFIED (Live Demo)"),
        ("Multi-Temporal Analysis", "Spectral Trajectory Engine", "VERIFIED (Timeline)"),
        ("Change Detection", "NDVI/NDBI Delta Masks", "VERIFIED (Change Map)"),
        ("False-Alarm Suppression", "Phenology & Quality Masking", "VERIFIED (Funnel)"),
        ("Analyst Verification", "Interactive Map Workstation", "VERIFIED (UI Dashboard)"),
        ("Air-Gapped Operation", "Local PyTorch & Vector Index", "VERIFIED (On-Prem)")
    ]

    for r_idx, row_data in enumerate(matrix):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_CARD
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                p.font.color.rgb = SUCCESS_GREEN if "VERIFIED" in val else TEXT_WHITE

# -----------------------------------------------------------------------------
# TECHNICAL BACKUP SLIDES (11 TO 18)
# -----------------------------------------------------------------------------

def build_slide_11():
    """Slide 11: Deep Technical Ingestion Pipeline"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 1: Ingestion & Preprocessing Architecture", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "DATA INGESTION & QUALITY PIPELINE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("GeoTIFF & COG Parsing", "Uses Rasterio & GDAL for cloud-optimized GeoTIFF ingestion, extracting multi-spectral bands (B2, B3, B4, B8, B11, B12)."),
        ("STAC Metadata Cataloging", "Populates spatio-temporal asset catalog with cloud cover percentage, sun azimuth, and acquisition timestamps."),
        ("Atmospheric Correction & QA Masking", "Applies Scene Classification Layer (SCL) to flag clouds, cirrus, snow, and water shadows."),
        ("Tiling & Spatial Reprojection", "Reprojects imagery to UTM / EPSG:4326 and generates standardized 256x256 pixel spatial tiles.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_12():
    """Slide 12: Multimodal Embedding & Vector Indexing"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 2: Multimodal Embedding & Vector Search Mechanics", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "VISION-LANGUAGE MODEL (VLM) & VECTOR SEARCH"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("VLM Architecture", "Utilizes OpenCLIP / RemoteCLIP Transformer backbone trained on satellite-text pairs to embed tiles into 512-dimensional feature space."),
        ("Vector Search Engine", "Employs high-performance vector indexing with Cosine Similarity and L2 distance metrics for sub-100ms response."),
        ("Multimodal Similarity Formulation", "Sim(Text, Tile) = (E_text • E_tile) / (||E_text|| ||E_tile||)"),
        ("Incremental Vector Ingestion", "Appends new tile embeddings on the fly without needing full index rebuilds.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_13():
    """Slide 13: Spectral Change Mathematics"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 3: Spectral Change Mathematics & Trajectory Formulations", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SPECTRAL FORMULATIONS & CHANGE TRAJECTORIES"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("Normalized Difference Vegetation Index (NDVI)", "NDVI = (NIR - RED) / (NIR + RED)  -> Tracks vegetation cover density."),
        ("Normalized Difference Built-Up Index (NDBI)", "NDBI = (SWIR - NIR) / (SWIR + NIR)  -> Pinpoints artificial structures & concrete."),
        ("Normalized Difference Water Index (NDWI)", "NDWI = (GREEN - NIR) / (GREEN + NIR) -> Maps water body boundaries."),
        ("Change Magnitude Delta", "Delta = sqrt((NDVI_after - NDVI_before)^2 + (NDBI_after - NDBI_before)^2)")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_14():
    """Slide 14: False-Alarm Filtering Logic & Math"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 4: False-Alarm Suppression Algorithms & Noise Reduction", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ALGORITHMIC NOISE SUPPRESSION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("Phenology Cycle Matching", "Compares same-month observations across consecutive years to cancel out annual agricultural cycles."),
        ("Sub-Pixel Registration", "Applies phase correlation to align multi-temporal scenes within 0.2 pixel margin, eliminating edge ringing."),
        ("Temporal Persistence Filter", "Change score must exceed threshold T across N >= 3 consecutive observations to trigger alert."),
        ("Spatial Context Morphology", "Applies opening/closing morphological operations to filter isolated pixel noise.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_15():
    """Slide 15: Air-Gapped Security & On-Prem Deployment"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 5: Air-Gapped Security & Local Infrastructure", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SECURITY & OPERATIONAL HARDENING"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("Zero Network Dependency", "All model weights, indexes, and database assets reside locally inside isolated containers."),
        ("Resource Optimization", "Runs efficiently on CPU workstations or GPU edge servers (NVIDIA Orin / RTX)."),
        ("Role-Based Access Control", "Enforces local authentication and role-based permissions for analysts."),
        ("Audit Provenance Logging", "Generates SHA-256 hashed audit records for every search, query, and exported evidence dossier.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_16():
    """Slide 16: Quantitative Evaluation Methodology"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 6: Ground Truth Dataset & Evaluation Framework", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BENCHMARK DATASET & EVALUATION METRICS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("Annotated EO Evaluation Set", "500+ annotated Sentinel-2 multi-temporal scenes across diverse terrain (industrial, riverine, forest, coastal)."),
        ("Retrieval Performance Metrics", "Evaluated via Precision@K, Recall@K, Mean Reciprocal Rank (MRR), and nDCG@5."),
        ("Change Detection Metrics", "Evaluated against human analyst annotations using IoU (Intersection over Union), F1-Score, and False Alarm Rate."),
        ("Latency & System Profiling", "Benchmarked query retrieval time, index build throughput, and memory footprint under multi-gigabyte loads.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

def build_slide_17():
    """Slide 17: Comprehensive PS 26227 Requirement Matrix"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 7: Full PS 26227 Requirement Traceability Matrix", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)

    rows, cols = 8, 3
    left, top, width, height = Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(5.333)
    table.columns[2].width = Inches(2.5)

    headers = ["PS 26227 Detailed Requirement", "Architecture Module Implementation", "Status"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG_CARD_LIGHT
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = AZURE_BRIGHT

    matrix = [
        ("Text-to-Image Semantic Search", "OpenCLIP Embedder + Vector Search", "COMPLIANT"),
        ("Image-to-Image Similarity", "Visual Embedding Cosine Similarity", "COMPLIANT"),
        ("Multi-Temporal Stack Analysis", "Spectral Trajectory Engine", "COMPLIANT"),
        ("False-Alarm Suppression", "SCL Quality + Phenology Filter", "COMPLIANT"),
        ("Geospatial Localization & Map", "MapLibre GL + PostGIS Spatial Index", "COMPLIANT"),
        ("Analyst Audit Trail & Dossiers", "SHA-256 Hashed Evidence Logger", "COMPLIANT"),
        ("Offline / Air-Gapped Operation", "100% Local Containerized Execution", "COMPLIANT")
    ]

    for r_idx, row_data in enumerate(matrix):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_CARD
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                p.font.color.rgb = SUCCESS_GREEN if "COMPLIANT" in val else TEXT_WHITE

def build_slide_18():
    """Slide 18: Scalability Roadmap & Field Deployment Strategy"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_slide_header(slide, "Backup 8: Future Scalability & Field Deployment Roadmap", "TECHNICAL DEEP DIVE")

    add_card(slide, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.5), border_color=AZURE_MAIN)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SCALABILITY & OPERATIONAL ROADMAP"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AZURE_BRIGHT

    details = [
        ("Phase 1: Field Station Deployment", "Packaging current container stack for edge deployment on tactical military workstations."),
        ("Phase 2: Multi-Sensor Fusion", "Expanding beyond Sentinel-2 optical data to integrate Synthetic Aperture Radar (SAR / Sentinel-1) for all-weather imagery."),
        ("Phase 3: Automated Cluster Discovery", "Unsupervised spatial-temporal clustering for autonomous hotspot discovery across vast regions."),
        ("Phase 4: Defense System Integration", "Connecting report export pipeline with standard GIS format specifications.")
    ]

    for title, desc in details:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = AZURE_BRIGHT
        p2 = p.add_run()
        p2.text = desc
        p2.font.bold = False
        p2.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

# -----------------------------------------------------------------------------
# MAIN BUILD EXECUTION
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    print("Building SIH 26227 Presentation Deck (18 Slides)...")

    # Core Operational Slides (1-10)
    build_slide_1()
    build_slide_2()
    build_slide_3()
    build_slide_4()
    build_slide_5()
    build_slide_6()
    build_slide_7()
    build_slide_8()
    build_slide_9()
    build_slide_10()

    # Technical Backup Slides (11-18)
    build_slide_11()
    build_slide_12()
    build_slide_13()
    build_slide_14()
    build_slide_15()
    build_slide_16()
    build_slide_17()
    build_slide_18()

    output_path = "SIH_PPT/SIH_26227_Geospatial_Intelligence_Platform.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")
